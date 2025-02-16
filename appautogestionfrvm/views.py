import json
from django.shortcuts import render, redirect, get_object_or_404
from django.http import HttpResponse
from .models import Institucional, Curso, Alumno, Asistencia, Concepto
from .forms import CursoForm, AlumnoForm
from django.contrib.auth.views import LoginView
from django.views.decorators.csrf import csrf_exempt
from django.http import JsonResponse
from .forms import AlumnoForm
from django.contrib.auth.decorators import login_required
from .forms import CustomLoginForm
from django.contrib.auth.views import LoginView
from datetime import datetime
from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle, Image, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.units import inch
from reportlab.lib.pagesizes import letter, landscape
from reportlab.pdfgen import canvas
from django.utils import timezone
from django.contrib import messages
import mercadopago
from django.http import JsonResponse
from django.shortcuts import render
from django.shortcuts import render, get_object_or_404, redirect
from .models import Alumno, Concepto
from .forms import AlumnoForm
from .models import Alumno, Concepto, Pago
from django.views import View
from .models import Pago
from django.shortcuts import render, get_object_or_404
from .models import Alumno, Pago
from django.contrib.auth.mixins import LoginRequiredMixin, UserPassesTestMixin
from django.views.generic import TemplateView
from reportlab.lib.pagesizes import letter
from django.http import HttpResponse
from django.shortcuts import render, get_object_or_404
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter
from io import BytesIO
import PyPDF2
from .models import Curso, Alumno
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
from io import BytesIO
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, render
import PyPDF2
import openpyxl
from django.http import HttpResponse
from django.shortcuts import render
from .models import Curso, Asistencia
from django.contrib import messages
from io import BytesIO
from django.shortcuts import render, get_object_or_404
from django.http import HttpResponse
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib import colors
from reportlab.pdfgen import canvas
import PyPDF2
from .models import Curso, Alumno
from datetime import datetime
from django.shortcuts import render
from django.db.models import Q
from .models import Alumno, Concepto
from django.shortcuts import redirect
from django.contrib.auth.views import LoginView
from django.shortcuts import redirect
from .forms import CustomLoginForm
from django.http import HttpResponse
from django.template.loader import get_template
from xhtml2pdf import pisa
from io import BytesIO
from django.http import JsonResponse

def generar_link_de_pago(monto, email, concepto_descripcion):
    sdk = mercadopago.SDK("APP_USR-568258915695172-092413-e54eb6e8a0352062fa668357e0b40154-1676314633")

    preference_data = {
        "items": [
            {
                "title": concepto_descripcion,
                "quantity": 1,
                "unit_price": monto,
            }
        ],
        "payer": {
            "email": email
        },
        "auto_return": "approved"
    }

    # Crear la preferencia
    preference_response = sdk.preference().create(preference_data)
    preference = preference_response["response"]

    # Devolver el link de pago
    return preference["init_point"]


@login_required
def consultar_asistencias(request):
    cursos = Curso.objects.all()
    if request.method == 'POST':
        curso_id = request.POST.get('curso_id')
        fecha = request.POST.get('fecha')

        if curso_id and fecha:
            try:
                asistencias = Asistencia.objects.filter(curso_id=curso_id, fecha=fecha)
                if asistencias.exists():
                    return generate_excel(curso_id, fecha, asistencias)
                else:
                    messages.warning(request, "No se encontraron asistencias para esta fecha.")
            except Exception as e:
                messages.error(request, str(e))
        else:
            messages.error(request, "Por favor, seleccione un curso y una fecha.")

    return render(request, 'consultar_asistencias.html', {
        'cursos': cursos,
    })

def generate_excel(curso_id, fecha, asistencias):
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "Asistencias"

    headers = ["Nombre", "Apellido", "Tipo de Documento", "Documento", "Email", "Presente"]
    worksheet.append(headers)

    for asistencia in asistencias:
        worksheet.append([
            asistencia.alumno.nombre,
            asistencia.alumno.apellido,
            asistencia.alumno.tipo_documento,
            asistencia.alumno.numero_documento,
            asistencia.alumno.email,
            "Sí" if asistencia.presente else "No"
        ])

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    filename = f"Asistencias_Curso_{curso_id}Fecha{fecha}.xlsx"
    response['Content-Disposition'] = f'attachment; filename="{filename}"'

    workbook.save(response)

    return response

@login_required
def emitir_certificado(request):
    cursos = Curso.objects.all()
    if request.method == 'POST':
        curso_id = request.POST.get('curso')
        numero_documento = request.POST.get('numero_documento')
        fecha_resolucion = request.POST.get('fecha_resolucion')
        numero_resolucion = request.POST.get('numero_resolucion')
        curso = get_object_or_404(Curso, id=curso_id)
        alumno = get_object_or_404(Alumno, curso=curso, numero_documento=numero_documento)
        fecha_resolucion_obj = datetime.strptime(fecha_resolucion, '%Y-%m-%d') 
        meses_espanol = {
            1: 'enero',
            2: 'febrero',
            3: 'marzo',
            4: 'abril',
            5: 'mayo',
            6: 'junio',
            7: 'julio',
            8: 'agosto',
            9: 'septiembre',
            10: 'octubre',
            11: 'noviembre',
            12: 'diciembre',
        }
        mes_resolucion = meses_espanol.get(fecha_resolucion_obj.month, 'Mes desconocido')
        buffer = BytesIO()
        c = canvas.Canvas(buffer, pagesize=landscape(A4))

        margen_izquierdo = 150
        margen_superior = 400  

        c.setFont("Times-Roman", 25)
        c.drawString(margen_izquierdo, margen_superior, "El alumno     ")

        c.setFont("Times-BoldItalic", 25)
        nombre_completo = f"   {alumno.nombre} {alumno.apellido}, DNI N° {alumno.numero_documento}"
        c.drawString(margen_izquierdo + 95, margen_superior, nombre_completo)

        c.setFont("Times-Roman", 25)
        texto_dni_y_curso = f"ha satisfecho las condiciones exigidas por el curso de "

        texto_inicio_x = margen_izquierdo + 90 + c.stringWidth(nombre_completo, "Times-BoldItalic", 20) + 10
        c.drawString(150, margen_superior-30, texto_dni_y_curso)
        c.drawString(150,margen_superior-60,f"{alumno.curso}, dictado por la Institución ")
        c.drawString(150,margen_superior-90,"EducApp.")
        c.drawString(150, margen_superior-150, f"Con fecha de resolución el {fecha_resolucion_obj.day} de {mes_resolucion} del {fecha_resolucion_obj.year} y número")
        c.drawString(150,margen_superior-180,f"de resolución {numero_resolucion}.")
        c.showPage()
        c.save()
        buffer.seek(0)

        with open('staticfiles/pdfs/certificado_plantilla.pdf', 'rb') as template_file:
            template_pdf = PyPDF2.PdfReader(template_file)
            output_pdf = PyPDF2.PdfWriter()

            page = template_pdf.pages[0]
            overlay_pdf = PyPDF2.PdfReader(buffer)
            overlay_page = overlay_pdf.pages[0]
            page.merge_page(overlay_page)
            output_pdf.add_page(page)

            for i in range(1, len(template_pdf.pages)):
                output_pdf.add_page(template_pdf.pages[i])

            response = HttpResponse(content_type='application/pdf')
            response['Content-Disposition'] = f'attachment; filename="certificado_{alumno.nombre}.pdf"'
            output_pdf.write(response)
            return response

    return render(request, 'documentos/emitir_certificado.html', {'cursos': cursos})

class AdminDashboardView(LoginRequiredMixin, UserPassesTestMixin, TemplateView):
    template_name = 'paginas/admin_dashboard.html'
    login_url = '/login/'

    def test_func(self):
        return self.request.user.groups.filter(name='Administradores').exists()

@login_required
def confirmar_pago(request):
    if request.method == 'GET':
        payment_status = request.GET.get('status') 
        payment_id = request.GET.get('payment_id')  
        concepto_id = request.GET.get('concepto_id')
        email = request.GET.get('email')
        alumno = Alumno.objects.get(email=email)
        concepto = Concepto.objects.get(id=concepto_id)
        if payment_status == 'approved':
            Pago.objects.create(
                alumno=alumno,
                concepto=concepto,
                monto=concepto.monto,
                id_transaccion=payment_id,
                estado_pago='completado'
            )
        elif payment_status == 'pending':
            Pago.objects.create(
                alumno=alumno,
                concepto=concepto,
                monto=concepto.monto,
                id_transaccion=payment_id,
                estado_pago='pendiente'
            )
        return render(request, 'confirmacion_pago.html', {'status': payment_status, 'payment_id': payment_id})

@csrf_exempt
def webhook_pago(request):
    if request.method == 'POST':
        data = json.loads(request.body)
        payment_id = data.get('data', {}).get('id')
        status = data.get('data', {}).get('status')
        try:
            pago = Pago.objects.get(payment_id=payment_id)
            pago.status = status
            pago.save()
        except Pago.DoesNotExist:
            pass
            
        return JsonResponse({'status': 'success'}, status=200)
    return JsonResponse({'status': 'error'}, status=400)

class NotificacionPago(View):
    def post(self, request, *args, **kwargs):
        data = request.POST
        alumno = Alumno.objects.get(numero_documento=data['numero_documento'])
        concepto = Concepto.objects.get(id=data['concepto_id'])
        Pago.objects.update_or_create(
            alumno=alumno,
            concepto=concepto,
            defaults={'estado_pago': data['estado_pago'] == 'approved'}
        )
        return JsonResponse({'status': 'success'})
    
@login_required
def generar_cuota(request):
    query = request.GET.get('query', '') 
    alumnos = Alumno.objects.all()
    if query:
        alumnos = alumnos.filter(
            Q(nombre__icontains=query) | 
            Q(apellido__icontains=query) | 
            Q(numero_documento__icontains=query)
        )
    alumnos_conceptos_disponibles = {}
    for alumno in alumnos:
        conceptos_disponibles = Concepto.objects.filter(curso_id=alumno.curso.id)  
        alumnos_conceptos_disponibles[alumno] = conceptos_disponibles
    return render(request, 'generar_cuota.html', {'alumnos_conceptos_disponibles': alumnos_conceptos_disponibles})

@login_required
def generar_cuota_procesar(request):
    if request.method == 'POST':
        numero_documento = request.POST['numero_documento']
        codigo_concepto = request.POST['codigo_concepto']
        curso_id = request.POST['curso_id']
        alumno = get_object_or_404(Alumno, numero_documento=numero_documento)
        concepto = get_object_or_404(Concepto, codigo=codigo_concepto, curso_id=curso_id)
        Pago.objects.create(alumno=alumno, concepto=concepto)
        return redirect('listado_alumnos')

@login_required
def gestion_alumno(request, numero_documento):

    alumno_existente = Alumno.objects.filter(numero_documento=numero_documento).first()

    if request.method == 'POST':
        form = AlumnoForm(request.POST, instance=alumno_existente)
        if form.is_valid():
            alumno = form.save(commit=False)
            alumno.numero_documento = numero_documento
            
            concepto_id = request.POST.get('concepto')
            if concepto_id:
                alumno.concepto_id = concepto_id
            alumno.save()
            return redirect('confirmacion_inscripcion', numero_documento=numero_documento)
    else:
        form = AlumnoForm(instance=alumno_existente) if alumno_existente else AlumnoForm(initial={'numero_documento': numero_documento})
        form.fields['numero_documento'].widget.attrs['readonly'] = True

    conceptos = Concepto.objects.all()

    return render(request, 'alumnos/crear_alumno.html', {
        'form': form,
        'numero_documento': numero_documento,
        'alumno': alumno_existente,
        'conceptos': conceptos
    })

@login_required
def administrar_alumno(request):
    alumnos = Alumno.objects.all()
    conceptos = Concepto.objects.all()

    return render(request, 'administrar_alumno.html', {
        'alumnos': alumnos,
        'conceptos': conceptos,
    })

@login_required
def obtener_conceptos(request, curso_id):
    print("Request Method:", request.method) 
    print("Headers:", request.headers) 

    if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
        conceptos = Concepto.objects.filter(curso_id=curso_id).values('id', 'descripcion', 'monto')
        return JsonResponse(list(conceptos), safe=False)
    return JsonResponse({'error': 'No es una solicitud AJAX'}, status=400)

def checkout(request):
    if request.method == 'POST':
        data = json.loads(request.body) 
        amount = data.get('amount') 
        email = data.get('email') 
        if not amount:
            return JsonResponse({'error': 'Monto no proporcionado'}, status=400)
        sdk = mercadopago.SDK("APP_USR-568258915695172-092413-e54eb6e8a0352062fa668357e0b40154-1676314633")
        preference_data = {
            "items": [
                {
                    "title": "Inscripción al curso",
                    "quantity": 1,
                    "unit_price": float(amount)
                }
            ],
            "payer": {
                "email": email
            },
            "back_urls": {
                "success": "http://localhost:8000/pago_exitoso/",
                "failure": "http://localhost:8000/pago_fallido/",
                "pending": "http://localhost:8000/pago_pendiente/"
            },
            "auto_return": "approved",
        }
        preference = sdk.preference().create(preference_data)
        return JsonResponse({'preference_id': preference['response']['id']})
    return JsonResponse({'error': 'Método no permitido'}, status=405)

from django.http import HttpResponse
from django.shortcuts import get_object_or_404, render
from io import BytesIO
from xhtml2pdf import pisa
from django.template.loader import get_template
from .models import Alumno, Pago

def pago_exitoso(request):
    return render(request, 'pago_exitoso.html')

def descargar_comprobante_pago(request, numero_documento):
    alumno = get_object_or_404(Alumno, numero_documento=numero_documento)
    pago = get_object_or_404(Pago, alumno=alumno)

    template = get_template('comprobante_pago_pdf.html')  # Plantilla para el PDF
    context = {
        'alumno': alumno,
        'pago': pago,
    }
    html = template.render(context)
    result = BytesIO()
    pdf = pisa.pisaDocument(BytesIO(html.encode('UTF-8')), result)

    if not pdf.err:
        response = HttpResponse(result.getvalue(), content_type='application/pdf')
        response['Content-Disposition'] = f'attachment; filename="comprobante_pago_{alumno.numero_documento}.pdf"'
        return response
    return HttpResponse("Error al generar el PDF", status=500)


def pago_fallido(request):
    return render(request, 'pago_fallido.html')

def pago_pendiente(request):
    return render(request, 'pago_pendiente.html')


def inicio(request):
    return render(request, 'paginas/inicio.html')

def institucional(request):
    return render(request, 'paginas/institucional.html')

def cursos(request):
    query = request.GET.get('q', '')
    if query:
        cursos = Curso.objects.filter(titulo__icontains=query)
    else:
        cursos = Curso.objects.all()
    return render(request, 'cursos/index.html', {'cursos': cursos})

@login_required
def crear(request):
    formulario = CursoForm(request.POST or None,request.FILES or None)
    if formulario.is_valid():
        formulario.save()
        return redirect('cursos')    
    return render(request, 'cursos/crear.html', {'formulario': formulario})

@login_required
def editar(request,id):
    curso = Curso.objects.get(id=id)
    formulario = CursoForm(request.POST or None, request.FILES or None, instance=curso)
    if formulario.is_valid() and request.POST:
        formulario.save()
        return redirect('cursos')
    return render(request, 'cursos/editar.html',{'formulario': formulario})

@login_required
def eliminar(request,id):
    curso = Curso.objects.get(id=id)
    curso.delete()
    return redirect('cursos')

def institucional_view(request):
    apartados = Institucional.objects.all()
    return render(request, 'institucional.html', {'apartados': apartados})

class CustomLoginView(LoginView):
    form_class = CustomLoginForm
    template_name = 'paginas/login.html'

    def form_valid(self, form):
        user = form.get_user() 
        if user.groups.filter(name='Administradores').exists():
            return redirect('admin_dashboard')
        else:
            return redirect('home')
    def form_invalid(self, form):
        return super().form_invalid(form)

def contacto(request):
    formulario = CursoForm(request.POST or None,request.FILES or None)
    if formulario.is_valid():
        formulario.save()
        return redirect('inicio')    
    return render(request, 'paginas/contacto.html', {'formulario': formulario})

def lista_cursos(request):
    cursos = Curso.objects.all().order_by('id')
    query = request.GET.get('q')
    if query:
        cursos = cursos.filter(nombre__icontains=query)
    return render(request, 'cursos/index.html', {'cursos': cursos})

@login_required
def listado_alumnos_por_curso(request):
    cursos = Curso.objects.all()
    curso_seleccionado = request.GET.get('curso')
    curso_obj = None
    alumnos = []

    if curso_seleccionado:
        try:
            curso_obj = Curso.objects.get(id=curso_seleccionado)
            alumnos = Alumno.objects.filter(curso=curso_obj)
        except Curso.DoesNotExist:
            curso_obj = None

    return render(request, 'listado_alumnos_por_curso.html', {
        'cursos': cursos,
        'alumnos': alumnos,
        'curso_obj': curso_obj,
    })

@login_required
def buscar_alumno(request):
    if request.method == 'POST':
        numero_documento = request.POST.get('dniAlumno')
        alumnos = Alumno.objects.filter(numero_documento=numero_documento)
        if alumnos.exists():
            if alumnos.count() == 1:
                return redirect('gestion_alumno', numero_documento=numero_documento)
            else:
                return render(request, 'alumnos/lista_alumnos.html', {
                    'alumnos': alumnos,
                    'error': 'Se encontraron varios alumnos con el mismo número de documento.'
                })
        else:
            return redirect('gestion_alumno', numero_documento=numero_documento)
    return render(request, 'alumnos/buscar_alumno.html')

@login_required
def generar_pdf(request, numero_documento):
    alumno = get_object_or_404(Alumno, numero_documento=numero_documento)
    template = get_template('template_pdf.html')
    context = {
        'alumno': alumno,
        'curso': alumno.curso.nombre if alumno.curso else 'N/A',
    }
    html = template.render(context)
    result = BytesIO()
    pdf = pisa.pisaDocument(BytesIO(html.encode('UTF-8')), result)

    if not pdf.err:
        response = HttpResponse(result.getvalue(), content_type='application/pdf')
        response['Content-Disposition'] = f'attachment; filename="alumno_{alumno.numero_documento}.pdf"'
        return response
    return HttpResponse("Error al generar el PDF", status=500)


def chatbot_response(request):
    faq = {
        "Fotografia": "El curso de Fotografía tiene una duración de 10 clases y se cursa los jueves de 18:00 a 22:00 hs.",
        "fotografia": "El curso de Fotografía tiene una duración de 10 clases y se cursa los jueves de 18:00 a 22:00 hs.",
        "Marketing digital": "El curso de Marketing Digital tiene una duración de 3 meses y se cursa los lunes de 18:00 a 20:00 hs.",
        "marketing": "El curso de Marketing Digital tiene una duración de 3 meses y se cursa los lunes de 18:00 a 20:00 hs.",
        "Python": "El curso de Python tiene una duración de 3 meses y se cursa los miércoles de 19:00 a 22:00 hs.",
        "python": "El curso de Python tiene una duración de 3 meses y se cursa los miércoles de 19:00 a 22:00 hs.",
        "Oratoria": "El curso de Oratoria tiene una duración de 4 clases y se cursa los lunes de 19:00 a 21:00 hs.",
        "oratoria": "El curso de Oratoria tiene una duración de 4 clases y se cursa los lunes de 19:00 a 21:00 hs.",
        "Inscripcion": "Para inscripciones comunicate a nuestro whatsapp 3534788248",
        "Hola": "Hola, ¿En qué podemos ayudarte?",
        "hola": "Hola, ¿En qué podemos ayudarte?",
        "Holaa": "Hola, ¿En qué podemos ayudarte?",
    }
    user_message = request.GET.get('message', '')
    response = faq.get(user_message, "Lo siento, no tengo información sobre eso. Por favor, intenta con otra pregunta.")
    return JsonResponse({'response': response})

def faq_query(query):
    faq = {
        "Fotografia": "El curso de Fotografía tiene una duración de 10 clases y se cursa los jueves de 18:00 a 22:00 hs.",
        "fotografia": "El curso de Fotografía tiene una duración de 10 clases y se cursa los jueves de 18:00 a 22:00 hs.",
        "Marketing digital": "El curso de Marketing Digital tiene una duración de 3 meses y se cursa los lunes de 18:00 a 20:00 hs.",
        "marketing": "El curso de Marketing Digital tiene una duración de 3 meses y se cursa los lunes de 18:00 a 20:00 hs.",
        "Python": "El curso de Python tiene una duración de 3 meses y se cursa los miércoles de 19:00 a 22:00 hs.",
        "python": "El curso de Python tiene una duración de 3 meses y se cursa los miércoles de 19:00 a 22:00 hs.",
        "Oratoria": "El curso de Oratoria tiene una duración de 4 clases y se cursa los lunes de 19:00 a 21:00 hs.",
        "oratoria": "El curso de Oratoria tiene una duración de 4 clases y se cursa los lunes de 19:00 a 21:00 hs.",
        "Inscripcion": "Para inscripciones comunicate a nuestro whatsapp 3534788248",
        "Hola": "Hola, ¿En qué podemos ayudarte?",
        "hola": "Hola, ¿En qué podemos ayudarte?",
        "Holaa": "Hola, ¿En qué podemos ayudarte?",
        # Agrega más preguntas y respuestas aquí
    }
    return faq.get(query.lower(), "Lo siento, no tengo información sobre eso. Por favor, intenta con otra pregunta.")


@login_required
def confirmacion_inscripcion(request, numero_documento):
    alumno = get_object_or_404(Alumno, numero_documento=numero_documento)
    payment_id = request.GET.get('payment_id')
    status = request.GET.get('status','pendiente')
    concepto_id = alumno.concepto_id
    curso_id = alumno.curso_id
    concepto = get_object_or_404(Concepto, id=concepto_id)
    monto_concepto = concepto.monto
    try:
        Pago.objects.create(
            alumno=alumno,
            concepto=concepto,
            fecha_pago=timezone.now(),
            monto=monto_concepto,
            id_transaccion=payment_id,
            estado_pago=status
        )
        print("Pago guardado exitosamente.")
    except Exception as e:
        print(f'Error al guardar el pago: {e}')

    return render(request, 'alumnos/confirmacion_inscripcion.html', {
        'alumno': alumno,
        'numero_documento': numero_documento,
        'monto_concepto': monto_concepto,
        'concepto_id': concepto_id,
        'curso_id': curso_id,
        'payment_id': payment_id,
        'status': status
    })

@login_required
def tomar_asistencia(request):
    cursos = Curso.objects.all()

    if request.method == 'POST':
        curso_id = request.POST.get('curso')
        fecha_str = request.POST.get('fecha')
        curso = get_object_or_404(Curso, id=curso_id)

        try:
            fecha = datetime.strptime(fecha_str, '%Y-%m-%d').date()
        except ValueError:
            messages.error(request, "Formato de fecha inválido.")
            return render(request, 'tomar_asistencia.html', {
                'curso': curso,
                'fecha': fecha_str,
                'cursos': cursos
            })

        if fecha > timezone.now().date():
            messages.error(request, "La fecha no puede ser posterior a la fecha actual.")
            return render(request, 'tomar_asistencia.html', {
                'curso': curso,
                'fecha': fecha_str,
                'cursos': cursos
            })

        alumnos = Alumno.objects.filter(curso=curso)

        if 'guardar_asistencia' in request.POST:
            response = HttpResponse(content_type='application/pdf')
            response['Content-Disposition'] = f'attachment; filename="asistencia_{curso.nombre}_{fecha}.pdf"'
            pdf = SimpleDocTemplate(response, pagesize=landscape(letter))
            elements = []
            image_path = "C:/Users/Usuario/seminariointegrador/static/imagenes/logoeducapp.png"
            img = Image(image_path)
            img.drawHeight = 0.5 * inch
            img.drawWidth = 5 * inch
            img.hAlign = 'LEFT'
            elements.append(img)
            elements.append(Spacer(1, 0.5 * inch))
            styles = getSampleStyleSheet()
            styles.add(ParagraphStyle(name='Center', alignment=TA_CENTER))
            title = f"Asistencia del curso {curso.nombre}"
            date_info = f"Fecha: {fecha}"
            elements.append(Paragraph(title, styles['Title']))
            elements.append(Paragraph(date_info, styles['Heading2']))
            data = [['Apellido', 'Nombre', 'Tipo Documento', 'Número Documento','Email', 'Asistencia']]
            for alumno in alumnos:
                presente = request.POST.get(f'presente_{alumno.id}', 'off') == 'on'
                Asistencia.objects.create(alumno=alumno, curso=curso, fecha=fecha, presente=presente)
                data.append([
                    alumno.apellido,
                    alumno.nombre,
                    alumno.tipo_documento,
                    alumno.numero_documento,
                    alumno.email,
                    'Presente' if presente else 'Ausente'
                ])
            table = Table(data, colWidths=[1.5 * inch, 1.5 * inch, 1.5 * inch, 1.5 * inch, 2.5 * inch, 1.5 * inch, 1.5 * inch, 1.5 * inch, 2.5 * inch, 1.5 * inch])
            table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.blueviolet),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.black),
                ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
                ('BACKGROUND', (0, 1), (-1, -1), colors.whitesmoke),
                ('GRID', (0, 0), (-1, -1), 1, colors.black),
            ]))
            elements.append(table)
            pdf.build(elements)
            return response
        return render(request, 'tomar_asistencia.html', {
            'curso': curso,
            'fecha': fecha_str,
            'alumnos': alumnos,
            'cursos': cursos
        })
    return render(request, 'tomar_asistencia.html', {'cursos': cursos})